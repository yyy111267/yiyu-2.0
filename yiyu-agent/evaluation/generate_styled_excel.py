#!/usr/bin/env python3
"""
生成美化版 Golden Case 主表 Excel（.xlsx）。
样式参考模板：深色标题栏 + 浅灰版本条 + 浅蓝表头 + 隔行变色。
字段全部使用中文名。JSON 字段做 pretty-print 便于阅读。
数据来源：golden_cases_25_v1.csv
"""
import csv
import json
import os
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, NamedStyle
)
from openpyxl.utils import get_column_letter

# ── 路径 ──
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CSV_PATH = os.path.join(BASE_DIR, "golden_cases_25_v1.csv")
XLSX_PATH = os.path.join(BASE_DIR, "golden_cases_25_v1_美化版.xlsx")

# ── 英文字段 → 中文名映射（顺序即列顺序）──
COLUMN_MAP = [
    ("case_id",              "Case ID",          22),
    ("name",                 "案例名称",          28),
    ("test_target",          "被测对象",          18),
    ("primary_set",          "内容性质",          10),
    ("regression_role",      "回归角色",          10),
    ("level",                "测试层级",          10),
    ("lifecycle",            "生命周期",          10),
    ("old_source",           "原始来源",          22),
    ("origin_type",          "来源类型",          16),
    ("owner",                "负责人",            10),
    ("fixture_id",           "数据快照ID",        26),
    ("fixture_version",      "快照版本",          10),
    ("input_message",        "输入",              36),
    ("input_metadata",       "输入元数据",        30),
    ("expected_route",       "预期路由",          24),
    ("expected_skill",       "预期技能",          18),
    ("required_tools",       "必调工具",          24),
    ("forbidden_tools",      "禁调工具",          24),
    ("tool_order",           "工具调用顺序",      24),
    ("expected_final_state", "预期终态",          32),
    ("field_assertions",     "字段断言",          40),
    ("hard_key_points",      "硬性要点",          36),
    ("soft_key_points",      "软性要点",          28),
    ("rubric_id",            "评分标准ID",        24),
    ("change_note",          "变更备注",          36),
]

JSON_FIELDS = {"input_metadata", "required_tools", "forbidden_tools", "tool_order", "field_assertions"}

# ── 配色（参考模板）──
COLOR_BANNER_BG   = "1B2A4A"   # 深色标题栏
COLOR_BANNER_FG   = "FFFFFF"    # 标题文字白
COLOR_VERSION_BG   = "E8EBF0"   # 浅灰版本条
COLOR_VERSION_FG   = "3A4A5C"   # 版本条文字
COLOR_HEADER_BG    = "D4E2F3"   # 浅蓝表头
COLOR_HEADER_FG    = "1B2A4A"   # 表头文字
COLOR_ROW_ALT      = "EAF2FB"   # 隔行浅蓝
COLOR_ROW_WHITE    = "FFFFFF"   # 白色行
COLOR_BORDER       = "B8C4D0"   # 边框浅灰蓝

# ── 样式对象 ──
thin_border = Border(
    left=Side(style="thin", color=COLOR_BORDER),
    right=Side(style="thin", color=COLOR_BORDER),
    top=Side(style="thin", color=COLOR_BORDER),
    bottom=Side(style="thin", color=COLOR_BORDER),
)

banner_font = Font(name="微软雅黑", size=16, bold=True, color=COLOR_BANNER_FG)
banner_fill = PatternFill(start_color=COLOR_BANNER_BG, end_color=COLOR_BANNER_BG, fill_type="solid")
banner_align = Alignment(horizontal="left", vertical="center", indent=1)

version_font = Font(name="微软雅黑", size=10, color=COLOR_VERSION_FG, italic=False)
version_fill = PatternFill(start_color=COLOR_VERSION_BG, end_color=COLOR_VERSION_BG, fill_type="solid")
version_align = Alignment(horizontal="left", vertical="center", indent=1)

header_font = Font(name="微软雅黑", size=10, bold=True, color=COLOR_HEADER_FG)
header_fill = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

cell_font = Font(name="微软雅黑", size=9, color="2A2A2A")
cell_align = Alignment(horizontal="left", vertical="top", wrap_text=True)
cell_align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

# 短字段居中显示
CENTER_FIELDS = {"test_target", "primary_set", "regression_role", "level", "lifecycle",
                  "origin_type", "owner", "fixture_version", "expected_skill", "rubric_id"}


def pretty_json(raw: str) -> str:
    """JSON 字符串 pretty-print；空值或解析失败返回原值。"""
    if not raw or not raw.strip():
        return raw
    try:
        obj = json.loads(raw)
        if isinstance(obj, (dict, list)) and obj:
            return json.dumps(obj, ensure_ascii=False, indent=2)
        return raw
    except (json.JSONDecodeError, TypeError):
        return raw


def load_cases():
    """从 CSV 读取数据。"""
    with open(CSV_PATH, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        return list(reader)


def build_main_sheet(ws, cases):
    """构建 Case 主表。"""
    n_cols = len(COLUMN_MAP)
    last_col_letter = get_column_letter(n_cols)

    # ── Row 1: 标题栏（合并）──
    ws.merge_cells(f"A1:{last_col_letter}1")
    c = ws["A1"]
    c.value = "以渔 Agent · Golden Case 主表（首批 25 条）"
    c.font = banner_font
    c.fill = banner_fill
    c.alignment = banner_align
    ws.row_dimensions[1].height = 38

    # ── Row 2: 版本信息条（合并）──
    ws.merge_cells(f"A2:{last_col_letter}2")
    c = ws["A2"]
    func_count = sum(1 for r in cases if r["primary_set"] == "functional")
    sec_count = sum(1 for r in cases if r["primary_set"] == "security")
    cand_count = sum(1 for r in cases if r["regression_role"] == "核心候选")
    sent_count = sum(1 for r in cases if r["regression_role"] == "哨兵")
    c.value = (
        f"数据集版本：golden-v1.0-20260829；共 {len(cases)} 条核心样本"
        f"（功能 {func_count} 条 / 安全 {sec_count} 条）；"
        f"回归候选 {cand_count} 条 + 哨兵 {sent_count} 条 = 回归集视图 {cand_count + sent_count} 条。"
        f"  全部 25 条录入同一张主表，按字段筛选视图，不分表。"
    )
    c.font = version_font
    c.fill = version_fill
    c.alignment = version_align
    ws.row_dimensions[2].height = 24

    # ── Row 3: 表头 ──
    for col_idx, (_, cn_name, _) in enumerate(COLUMN_MAP, start=1):
        c = ws.cell(row=3, column=col_idx, value=cn_name)
        c.font = header_font
        c.fill = header_fill
        c.alignment = header_align
        c.border = thin_border
    ws.row_dimensions[3].height = 30

    # ── Row 4+: 数据行 ──
    for row_idx, case in enumerate(cases, start=4):
        is_alt = (row_idx - 4) % 2 == 1
        row_fill = PatternFill(
            start_color=COLOR_ROW_ALT if is_alt else COLOR_ROW_WHITE,
            end_color=COLOR_ROW_ALT if is_alt else COLOR_ROW_WHITE,
            fill_type="solid",
        )
        for col_idx, (en_key, _, _) in enumerate(COLUMN_MAP, start=1):
            raw_val = case.get(en_key, "")
            # JSON 字段 pretty-print
            if en_key in JSON_FIELDS:
                val = pretty_json(raw_val)
            else:
                val = raw_val

            c = ws.cell(row=row_idx, column=col_idx, value=val if val != "" else None)
            c.font = cell_font
            c.fill = row_fill
            c.border = thin_border
            if en_key in CENTER_FIELDS:
                c.alignment = cell_align_center
            else:
                c.alignment = cell_align

        # 行高根据内容估算（JSON 字段行数）
        max_lines = 1
        for en_key, _, _ in COLUMN_MAP:
            if en_key in JSON_FIELDS:
                v = case.get(en_key, "")
                if v:
                    lines = v.count("\n") + pretty_json(v).count("\n") + 1
                    max_lines = max(max_lines, min(lines, 12))
            else:
                v = case.get(en_key, "")
                if v and len(v) > 40:
                    est = len(v) // 35 + 1
                    max_lines = max(max_lines, min(est, 8))
        ws.row_dimensions[row_idx].height = max(18, min(max_lines * 14, 160))

    # ── 列宽 ──
    for col_idx, (_, _, width) in enumerate(COLUMN_MAP, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # ── 冻结窗格（冻结前 3 行 + 前 2 列）──
    ws.freeze_panes = "C4"

    # ── 自动筛选 ──
    ws.auto_filter.ref = f"A3:{last_col_letter}{3 + len(cases)}"


def build_glossary_sheet(ws):
    """构建口径与说明 sheet。"""
    entries = [
        ("字段", "说明", "取值范围 / 示例"),
        ("Case ID", "案例唯一标识：2字母能力前缀 + 全局递增3位序号，如 RT-001、MT-004", "RT-001 ~ RS-025，全局唯一不重复"),
        ("案例名称", "简短描述被测能力", "—"),
        ("被测对象", "本条 case 主要测试的对象", "Agent / skill:route / skill:delivery-validation / skill:light-answer / 整链"),
        ("内容性质", "题目考正常能力还是考红线（正交轴1）", "functional / security"),
        ("回归角色", "题目处于什么生命周期状态（正交轴2）", "核心候选 / 哨兵 / 否"),
        ("测试层级", "stage=单技能单元测试；e2e=端到端全链路", "stage / e2e"),
        ("生命周期", "draft→active→frozen 状态流转", "draft / active / frozen"),
        ("原始来源", "现有文件与旧 ID，便于迁移对照", "01_intent: ir_a01 等"),
        ("来源类型", "existing_test=现有测试；historical_badcase=历史事故哨兵", "existing_test / historical_badcase"),
        ("负责人", "fixture 和 case 的负责人", "待补充"),
        ("数据快照ID", "待创建的固定数据快照 ID", "route-fixture-v1 等"),
        ("快照版本", "fixture 版本号", "v1"),
        ("输入", "用户输入消息", "—"),
        ("输入元数据", "JSON：附加输入条件（用户记忆、tier、fixture 数据等）", "JSON 对象"),
        ("预期路由", "预期的路由结果", "light_answer / research_task + deep-research / out_of_scope / 不适用"),
        ("预期技能", "预期处理该 case 的技能模块", "light-answer / deep-research / delivery-validation 等"),
        ("必调工具", "JSON 数组：必须调用的工具", "JSON 数组"),
        ("禁调工具", "JSON 数组：禁止调用的工具", "JSON 数组"),
        ("工具调用顺序", "JSON 数组：工具调用的预期顺序", "JSON 数组"),
        ("预期终态", "最终应达到的状态描述", "—"),
        ("字段断言", "JSON：机器可校验的字段级断言", "JSON 对象"),
        ("硬性要点", "Hard 断言，必须满足", "—"),
        ("软性要点", "Soft 断言，参与平均分", "—"),
        ("评分标准ID", "关联的 Rubric 编号", "R-ROUTE-RULE-001 等"),
        ("变更备注", "变更记录、待办事项", "—"),
        ("", "", ""),
        ("设计原则", "", ""),
        ("单表原则", "25 条全部进同一张 Case 主表，用字段标注归属，不拆功能/安全/回归表", ""),
        ("两条正交轴", "内容性质（primary_set）× 使用纪律（regression_role），不是三种并列的表", ""),
        ("回归集视图", "regression_role in (核心候选, 哨兵) 筛出的视图，不是独立复制的另一张表", ""),
        ("哨兵身份", "primary_set 仍为 functional；哨兵身份由 regression_role=哨兵 + origin_type=historical_badcase 表达", ""),
        ("JSON 字段", "field_assertions / required_tools / forbidden_tools / tool_order / input_metadata 均为 JSON 字符串，Runner 直接解析", ""),
        ("", "", ""),
        ("Case ID 前缀对照", "", ""),
        ("RT", "路由（Route）—— 意图识别、轻回答/研究路由、规则匹配", "RT-001 ~ RT-003, RT-023"),
        ("MT", "指标计算（Metric）—— ROE、FCF、ROIC 等财务指标", "MT-004 ~ MT-006"),
        ("DT", "数据取数（Data）—— 官方来源优先、快照、实体消歧", "DT-007 ~ DT-009"),
        ("RS", "研究（Research）—— 单标的研究、双标的对比、预算降级、Tool失败降级", "RS-010 ~ RS-012, RS-025"),
        ("MM", "记忆（Memory）—— 用户方法论边界", "MM-013"),
        ("LT", "轻回答（Light answer）—— 单点数据查询，不启动深度研究", "LT-014"),
        ("DV", "准出校验（Delivery validation）—— validate_conclusion 规则拦截", "DV-015 ~ DV-018, DV-024"),
        ("BD", "边界（Boundary）—— 拒绝短期涨跌预测等超范围请求", "BD-019"),
        ("IJ", "注入（Injection）—— Prompt 注入、系统内容泄露防护", "IJ-020"),
        ("TD", "交易（Trade）—— 自动交易请求越权拦截", "TD-021"),
        ("SR", "来源（Source）—— 不可信来源（论坛爆料）防护", "SR-022"),
    ]

    # 标题
    ws.merge_cells("A1:C1")
    c = ws["A1"]
    c.value = "口径与字段说明"
    c.font = banner_font
    c.fill = banner_fill
    c.alignment = banner_align
    ws.row_dimensions[1].height = 36

    for row_idx, (a, b, c_val) in enumerate(entries, start=2):
        is_header = (row_idx == 2) or (a == "设计原则") or (a == "Case ID 前缀对照")
        fill = header_fill if is_header else (
            PatternFill(start_color=COLOR_ROW_ALT, end_color=COLOR_ROW_ALT, fill_type="solid")
            if (row_idx % 2 == 0) else
            PatternFill(start_color=COLOR_ROW_WHITE, end_color=COLOR_ROW_WHITE, fill_type="solid")
        )
        for col_idx, val in enumerate([a, b, c_val], start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val if val else None)
            cell.font = Font(name="微软雅黑", size=10, bold=is_header, color="2A2A2A")
            cell.fill = fill
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = thin_border

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 55
    ws.column_dimensions["C"].width = 45
    ws.freeze_panes = "A3"


def main():
    cases = load_cases()
    print(f"Loaded {len(cases)} cases from CSV")

    wb = Workbook()

    # Sheet 1: Case 主表
    ws_main = wb.active
    ws_main.title = "Case主表"
    build_main_sheet(ws_main, cases)

    # Sheet 2: 口径与说明
    ws_glossary = wb.create_sheet("口径与说明")
    build_glossary_sheet(ws_glossary)

    wb.save(XLSX_PATH)
    print(f"Saved styled Excel to: {XLSX_PATH}")
    print(f"  Sheet 1: Case主表 ({len(cases)} rows x {len(COLUMN_MAP)} cols)")
    print(f"  Sheet 2: 口径与说明")


if __name__ == "__main__":
    main()
